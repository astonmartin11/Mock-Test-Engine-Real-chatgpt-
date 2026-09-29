from __future__ import annotations
from dataclasses import dataclass
from io import BytesIO
import fitz
from pptx import Presentation

@dataclass
class ParsedPage:
    page_number: int
    text: str
    has_images: bool = False
    has_tables: bool = False
    has_formulas: bool = False
    needs_vision: bool = False

def parse_pdf(data: bytes) -> list[ParsedPage]:
    doc = fitz.open(stream=data, filetype='pdf')
    pages=[]
    try:
        for i,page in enumerate(doc):
            text=page.get_text('text') or ''
            images=len(page.get_images(full=True))
            pages.append(ParsedPage(i+1,text,has_images=images>0,needs_vision=(len(text.strip())<150 and images>0)))
    finally:
        doc.close()
    return pages

def parse_pptx(data: bytes) -> list[ParsedPage]:
    prs=Presentation(BytesIO(data)); out=[]
    for i,slide in enumerate(prs.slides):
        texts=[]
        for shape in slide.shapes:
            if hasattr(shape,'text') and shape.text.strip(): texts.append(shape.text.strip())
        out.append(ParsedPage(i+1,'\n'.join(texts),has_images=any(getattr(s,'shape_type',None)==13 for s in slide.shapes)))
    return out

def parse_document(filename: str,mime_type: str,data: bytes) -> list[ParsedPage]:
    lower=filename.lower()
    if lower.endswith('.pdf') or mime_type=='application/pdf': return parse_pdf(data)
    if lower.endswith('.pptx'): return parse_pptx(data)
    if lower.endswith('.txt'): return [ParsedPage(1,data.decode('utf-8',errors='replace'))]
    raise ValueError('Unsupported document type. Use PDF, PPTX, or TXT.')
